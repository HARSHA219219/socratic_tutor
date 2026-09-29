"""Document uploads (PDF / DOCX / TXT): validation, extraction, storage, secure access, LLM payload."""
import asyncio
import io
from pathlib import Path
from types import SimpleNamespace

import pytest
from sqlalchemy import select

from app.core.config import settings
from app.db.database import AsyncSessionLocal
from app.models.domain import Message
from app.services import document_storage as ds
from app.services.llm_service import build_llm_messages

PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64


def make_pdf(text: str) -> bytes:
    """Minimal valid one-page PDF with a Helvetica text line."""
    return make_pdf_stream(f"BT /F1 18 Tf 20 60 Td ({text}) Tj ET".encode())


def make_pdf_stream(stream: bytes) -> bytes:
    """One-page PDF with the given content stream (e.g. vector shapes and no text)."""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 144] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return out


def make_docx(paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Cell A"
    table.rows[0].cells[1].text = "Cell B"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def new_session(client, headers):
    return client.post("/api/sessions/", json={}, headers=headers).json()["id"]


def upload(client, headers, sid, name, data, content_type="application/octet-stream", message=""):
    return client.post(
        "/api/chat/document",
        data={"session_id": sid, "message": message},
        files={"document": (name, data, content_type)},
        headers=headers,
    )


# ---------- extraction (unit) ----------

def test_extract_pdf_docx_txt():
    pdf_mime = ds.validate_document(make_pdf("Photosynthesis makes glucose"), "notes.pdf")
    assert pdf_mime == ds.PDF_MIME
    text, status = ds.extract_text(make_pdf("Photosynthesis makes glucose"), pdf_mime)
    assert "Photosynthesis makes glucose" in text and status == "ok"

    docx = make_docx(["Newton's second law", "F = ma"])
    assert ds.validate_document(docx, "law.docx") == ds.DOCX_MIME
    text, _ = ds.extract_text(docx, ds.DOCX_MIME)
    assert "Newton's second law" in text and "F = ma" in text and "Cell A | Cell B" in text

    txt = "Café ☕ résumé\nline two".encode("utf-8")
    assert ds.validate_document(txt, "notes.txt") == ds.TXT_MIME
    assert ds.extract_text(txt, ds.TXT_MIME)[0] == "Café ☕ résumé\nline two"


def test_type_comes_from_content_not_filename():
    # A PDF named .txt is still detected as PDF; text renamed to .pdf is rejected
    assert ds.validate_document(make_pdf("x"), "sneaky.txt") == ds.PDF_MIME
    with pytest.raises(ds.DocumentError):
        ds.validate_document(b"just some text", "fake.pdf")
    # A zip that isn't a Word document
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "hi")
    with pytest.raises(ds.DocumentError):
        ds.validate_document(buf.getvalue(), "archive.docx")
    # Binary data with a .txt name
    with pytest.raises(ds.DocumentError):
        ds.validate_document(b"\x00\x01\x02binary", "data.txt")


def test_text_truncated_to_limit(monkeypatch):
    monkeypatch.setattr(settings, "MAX_DOCUMENT_TEXT_CHARS", 50)
    text, status = ds.extract_text(("word " * 100).encode(), ds.TXT_MIME)
    assert len(text) == 50 and status == "truncated"


def test_sanitize_filename():
    assert ds.sanitize_filename("../../etc/passwd", ds.TXT_MIME) == "passwd.txt"
    assert ds.sanitize_filename("C:\\Users\\me\\Notes.PDF", ds.PDF_MIME) == "Notes.PDF"
    assert ds.sanitize_filename('a<b>:"c|?*.docx', ds.DOCX_MIME) == "abc.docx"
    assert ds.sanitize_filename("", ds.PDF_MIME) == "document.pdf"
    assert ds.sanitize_filename("...hidden", ds.TXT_MIME) == "hidden.txt"
    with pytest.raises(ValueError):
        ds.resolve_document_path("../secret.txt")
    with pytest.raises(ValueError):
        ds.resolve_document_path("notes.pdf")


# ---------- endpoint ----------

def test_document_upload_storage_llm_and_access(client, make_student, admin_headers, fake_llm):
    s = make_student()
    sid = new_session(client, s["headers"])
    pdf = make_pdf("Mitochondria is the powerhouse of the cell")
    res = upload(client, s["headers"], sid, "../../Biology Notes.pdf", pdf, "text/plain", message="Explain this")
    assert res.status_code == 200, res.text
    msg = res.json()["user_message"]
    assert msg["document_name"] == "Biology Notes.pdf"
    assert msg["document_mime_type"] == "application/pdf"
    assert msg["document_size_bytes"] == len(pdf)
    assert msg["document_text_status"] == "ok"
    assert "document_text" not in msg  # extracted text stays server-side

    # Stored on disk under a random name; SQLite has metadata + text, not the binary
    async def stored():
        async with AsyncSessionLocal() as db:
            return (await db.execute(select(Message).where(Message.id == msg["id"]))).scalars().first()
    row = asyncio.run(stored())
    assert row.document_key.endswith(".pdf") and "Biology" not in row.document_key
    assert (Path(settings.DOCUMENT_STORAGE_DIR) / row.document_key).read_bytes() == pdf
    assert "powerhouse" in row.document_text

    # The LLM got the student's text AND the extracted document text
    sent = fake_llm.calls[0]["messages"][-1]
    assert sent["role"] == "user" and sent["content"].startswith("Explain this")
    assert "powerhouse of the cell" in sent["content"] and "<document>" in sent["content"]

    # Survives reload
    history = client.get(f"/api/chat/session/{sid}", headers=s["headers"]).json()
    assert history[0]["document_name"] == "Biology Notes.pdf"

    # Secure download: owner + admin only
    url = f"/api/documents/{msg['id']}"
    own = client.get(url, headers=s["headers"])
    assert own.status_code == 200 and own.content == pdf
    assert own.headers["content-type"] == "application/pdf"
    assert "attachment" in own.headers["content-disposition"]
    assert own.headers["x-content-type-options"] == "nosniff"
    assert client.get(url).status_code == 401
    assert client.get(url, headers=make_student()["headers"]).status_code == 404
    assert client.get(url, headers=admin_headers).status_code == 200
    assert client.get("/api/documents/not-a-real-id", headers=admin_headers).status_code == 404

    admin_log = client.get(f"/api/admin/students/{s['student']['id']}/sessions/{sid}", headers=admin_headers).json()
    assert admin_log[0]["document_name"] == "Biology Notes.pdf"
    assert admin_log[0]["document_text_status"] == "ok"


def test_docx_and_txt_uploads_and_title(client, make_student, fake_llm):
    s = make_student()
    sid = new_session(client, s["headers"])
    res = upload(client, s["headers"], sid, "essay.docx", make_docx(["My thesis is about climate."]))
    assert res.status_code == 200, res.text
    assert res.json()["user_message"]["content"] == ""  # document-only message
    assert "My thesis is about climate." in fake_llm.calls[0]["messages"][-1]["content"]
    assert res.json()["session_title"]  # chat still gets a name

    res = upload(client, s["headers"], sid, "q.txt", b"What is entropy?")
    assert res.status_code == 200
    assert res.json()["user_message"]["document_mime_type"] == "text/plain"


def test_rejects_bad_documents_and_saves_nothing(client, make_student, monkeypatch):
    s = make_student()
    sid = new_session(client, s["headers"])
    cases = [
        ("empty.pdf", b"", 400, "empty"),
        ("photo.pdf", PNG, 400, "image button"),
        ("script.exe", b"MZ\x90\x00binary", 400, "Unsupported"),
        ("notes.pdf", b"not really a pdf", 400, "Unsupported"),
        ("blank.pdf", make_pdf(""), 422, "blank"),
        ("broken.pdf", b"%PDF-1.4\n garbage", 422, "PDF"),
    ]
    for name, data, code, words in cases:
        res = upload(client, s["headers"], sid, name, data)
        assert res.status_code == code, (name, res.text)
        assert words in res.json()["detail"], (name, res.json())

    monkeypatch.setattr(settings, "MAX_DOCUMENT_BYTES", 100)
    big = upload(client, s["headers"], sid, "big.txt", b"a" * 500)
    assert big.status_code == 413 and "too large" in big.json()["detail"]

    assert client.get(f"/api/chat/session/{sid}", headers=s["headers"]).json() == []


def test_document_upload_requires_owner_session(client, make_student):
    a, b = make_student(), make_student()
    sid = new_session(client, a["headers"])
    assert upload(client, b["headers"], sid, "q.txt", b"hi").status_code == 404
    assert client.post("/api/chat/document", data={"session_id": sid}, files={"document": ("q.txt", b"hi")}).status_code == 401


def test_llm_history_limits_old_documents():
    def m(i, role="user", doc=False, image=False):
        return SimpleNamespace(
            id=str(i), role=role, content=f"m{i}",
            image_key=f"{i}.png" if image else None, image_mime_type="image/png" if image else None,
            document_key=f"{i}.pdf" if doc else None, document_name=f"file{i}.pdf",
            document_text=f"TEXT{i}" if doc else None, document_text_status="ok",
        )

    async def loader(key, mime):
        return "data:image/png;base64,AAAA"

    history = [m(1, doc=True), m(2, "assistant"), m(3, doc=True), m(4, image=True), m(5, doc=True)]
    messages, has_images = asyncio.run(build_llm_messages("sys", history, loader))
    contents = [msg["content"] for msg in messages]
    assert "TEXT1" not in contents[1] and 'attached a document earlier: "file1.pdf"' in contents[1]
    assert "TEXT3" in contents[3] and "TEXT5" in contents[5]
    # Image flow unchanged
    assert has_images and contents[4][1]["type"] == "image_url"


# ---------- PDFs with no text layer (scans / "Print to PDF" outlined text) ----------

# Filled shapes but no text operators, like letters converted to vector outlines
OUTLINED_PDF = make_pdf_stream(b"0 0 0 rg 20 20 200 60 re f 0.2 0.4 0.8 rg 240 30 80 40 re f")


def test_textless_pdf_is_accepted_as_page_images():
    text, status = ds.extract_text(OUTLINED_PDF, ds.PDF_MIME)
    assert (text, status) == ("", "image")
    jpegs, blank = ds.render_pdf_pages(OUTLINED_PDF, 4)
    assert len(jpegs) == 1 and jpegs[0][:3] == b"\xff\xd8\xff" and not blank


def test_textless_pdf_pages_go_to_vision_model(client, make_student, admin_headers, fake_llm):
    s = make_student()
    sid = new_session(client, s["headers"])
    res = upload(client, s["headers"], sid, "printed.pdf", OUTLINED_PDF, message="Summarise this")
    assert res.status_code == 200, res.text
    assert res.json()["user_message"]["document_text_status"] == "image"

    last = fake_llm.calls[0]["messages"][-1]
    assert isinstance(last["content"], list)
    assert last["content"][0]["text"].startswith("Summarise this")
    assert "no selectable text" in last["content"][0]["text"]
    images = [p for p in last["content"] if p["type"] == "image_url"]
    assert images and images[0]["image_url"]["url"].startswith("data:image/jpeg;base64,")

    logs = client.get(f"/api/admin/students/{s['student']['id']}/sessions/{sid}", headers=admin_headers).json()
    assert logs[0]["document_text_status"] == "image"
